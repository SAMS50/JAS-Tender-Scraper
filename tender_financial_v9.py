
"""
JAS Tender Scraper — V9 Financial Evidence Audit

Reads:
  tender_metadata_batch_v8_1.csv
  downloaded_tenders/*.zip

Writes:
  tender_financial_candidates_v9.csv
  tender_financial_audit_v9.json

No automatic financial promotion.
No fabricated values.
No modifications to V5–V8.1 outputs.
"""

import csv
import io
import json
import re
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import ZipFile

import pymupdf
import xlrd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
ZIP_DIR = ROOT / "downloaded_tenders"

INPUT = ROOT / "tender_metadata_batch_v8_1.csv"
OUTPUT = ROOT / "tender_financial_candidates_v9.csv"
AUDIT = ROOT / "tender_financial_audit_v9.json"

MAX_PDF_PAGES = 30
MAX_SHEET_ROWS = 200
MAX_SHEET_COLUMNS = 30

FIELDS = [
    "zip_filename",
    "tender_id",
    "tender_id_status",
    "document",
    "location",
    "candidate_field",
    "candidate_value",
    "candidate_basis",
    "source_label",
    "evidence",
    "review_status",
]

# These patterns classify labels, not financial values.
LABELS = [
    (
        "estimated_value_inr",
        "explicit_estimated_value",
        re.compile(
            r"\b(?:estimated\s+(?:cost|value|amount)|"
            r"estimated\s+contract\s+value|"
            r"tender\s+value|"
            r"estimated\s+value\s+of\s+work|"
            r"estimated\s+cost\s+of\s+work)\b",
            re.I,
        ),
    ),
    (
        "emd_inr",
        "explicit_emd",
        re.compile(
            r"\b(?:earnest\s+money(?:\s+deposit)?|"
            r"e\.?\s*m\.?\s*d\.?|"
            r"bid\s+security)\b",
            re.I,
        ),
    ),
    (
        "tender_fee_inr",
        "explicit_tender_fee",
        re.compile(
            r"\b(?:tender\s+(?:document\s+)?fee|"
            r"cost\s+of\s+tender\s+document|"
            r"document\s+fee|"
            r"bid\s+document\s+fee)\b",
            re.I,
        ),
    ),
]

# Values with an explicit currency indicator are safest.
MONEY = re.compile(
    r"(?<!\w)(?:₹|Rs\.?|INR)\s*"
    r"([\d,]+(?:\.\d{1,2})?)"
    r"\s*(crores?|cr\.?|lakhs?|lacs?|thousands?)?",
    re.I,
)

# Bare amounts may be acceptable ONLY directly after
# a recognised label, e.g. "EMD: 25000".
BARE_MONEY = re.compile(
    r"^\s*[:=\-–]?\s*"
    r"([\d,]+(?:\.\d{1,2})?)"
    r"\s*(crores?|cr\.?|lakhs?|lacs?|thousands?)?"
    r"\s*(?:/-)?(?:\s|$)",
    re.I,
)

MULTIPLIERS = {
    "crore": Decimal("10000000"),
    "crores": Decimal("10000000"),
    "cr": Decimal("10000000"),
    "lakh": Decimal("100000"),
    "lakhs": Decimal("100000"),
    "lac": Decimal("100000"),
    "lacs": Decimal("100000"),
    "thousand": Decimal("1000"),
    "thousands": Decimal("1000"),
}

CONDITIONAL = re.compile(
    r"\b(?:exempt(?:ion|ed)?|waiv(?:e|ed|er)|"
    r"msme|mse|nsic|startup|"
    r"group[-\s]?[12]|"
    r"either|both\s+groups|"
    r"percentage|percent|%)\b",
    re.I,
)

NEGATION = re.compile(
    r"\b(?:not\s+applicable|nil|zero|"
    r"no\s+emd|no\s+tender\s+fee)\b",
    re.I,
)

GENERIC_DOCUMENTS = {
    "circular.pdf",
    "gcc2024.pdf",
}

SHARED_HASH_PREFIXES = {
    "253f3751960a772a",
    "03b049531944a625",
    "467343d85e9b47a5",
}


def normalize(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def parse_money(raw, unit=""):
    try:
        number = Decimal(str(raw).replace(",", ""))
        factor = MULTIPLIERS.get(
            normalize(unit).lower().rstrip("."),
            Decimal("1"),
        )
        result = number * factor
        if result < 0:
            return ""
        return format(result.quantize(Decimal("0.01")), "f")
    except (InvalidOperation, ValueError):
        return ""


def extract_amounts(fragment, allow_bare=False):
    values = []

    for match in MONEY.finditer(fragment):
        amount = parse_money(
            match.group(1),
            match.group(2) or "",
        )
        if amount:
            values.append(amount)

    if not values and allow_bare:
        match = BARE_MONEY.search(fragment)
        if match:
            amount = parse_money(
                match.group(1),
                match.group(2) or "",
            )
            if amount:
                values.append(amount)

    return list(dict.fromkeys(values))


def classify_evidence(text):
    if CONDITIONAL.search(text):
        return "CONDITIONAL_REVIEW"
    if NEGATION.search(text):
        return "NON_NUMERIC_REVIEW"
    return "NEEDS_REVIEW"


def scan_text(
    zip_name,
    base,
    document,
    location,
    text,
):
    """
    One candidate per explicitly labelled amount.

    Context is deliberately narrow to reduce accidental
    associations between unrelated financial fields.
    """
    candidates = []
    lines = [
        normalize(line)
        for line in text.splitlines()
        if normalize(line)
    ]

    for index, line in enumerate(lines):
        for field, basis, pattern in LABELS:
            matches = list(pattern.finditer(line))

            for label_match in matches:
                # Stop at another known financial label.
                fragment = line[label_match.end():]

                next_label_positions = [
                    other.start()
                    for _, _, other_pattern in LABELS
                    for other in other_pattern.finditer(fragment)
                ]

                if next_label_positions:
                    fragment = fragment[
                        :min(next_label_positions)
                    ]

                # Currency-marked values can appear after
                # the label. Bare values require adjacency.
                amounts = extract_amounts(
                    fragment,
                    allow_bare=True,
                )

                # If the label occupies its own line,
                # inspect the immediately following line.
                if not amounts and not fragment.strip(
                    " :-=\t"
                ) and index + 1 < len(lines):
                    next_line = lines[index + 1]

                    if not any(
                        p.search(next_line)
                        for _, _, p in LABELS
                    ):
                        amounts = extract_amounts(
                            next_line,
                            allow_bare=True,
                        )

                evidence = " | ".join(
                    lines[max(0, index - 1):index + 2]
                )[:600]

                status = classify_evidence(evidence)

                for amount in amounts:
                    candidates.append({
                        "zip_filename": zip_name,
                        "tender_id": base.get(
                            "tender_id", ""
                        ),
                        "tender_id_status": base.get(
                            "tender_id_status", ""
                        ),
                        "document": document,
                        "location": location,
                        "candidate_field": field,
                        "candidate_value": amount,
                        "candidate_basis": basis,
                        "source_label": label_match.group(),
                        "evidence": evidence,
                        "review_status": status,
                    })

    return candidates


def spreadsheet_rows(data, filename):
    if filename.lower().endswith(".xlsx"):
        wb = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
        )
        try:
            for sheet in wb.worksheets:
                for row_index, row in enumerate(
                    sheet.iter_rows(
                        max_row=MAX_SHEET_ROWS,
                        max_col=MAX_SHEET_COLUMNS,
                        values_only=True,
                    ),
                    start=1,
                ):
                    yield (
                        sheet.title,
                        row_index,
                        [normalize(cell) for cell in row],
                    )
        finally:
            wb.close()
    else:
        wb = xlrd.open_workbook(
            file_contents=data,
            on_demand=True,
        )
        try:
            for sheet in wb.sheets():
                for row_index in range(
                    min(sheet.nrows, MAX_SHEET_ROWS)
                ):
                    yield (
                        sheet.name,
                        row_index + 1,
                        [
                            normalize(
                                sheet.cell_value(
                                    row_index, col
                                )
                            )
                            for col in range(
                                min(
                                    sheet.ncols,
                                    MAX_SHEET_COLUMNS,
                                )
                            )
                        ],
                    )
        finally:
            wb.release_resources()


def inspect_zip(base):
    import hashlib

    name = base["zip_filename"]
    path = ZIP_DIR / name

    candidates = []
    errors = []
    scanned_pdfs = 0
    scanned_boqs = 0
    skipped_shared = 0

    if not path.exists():
        return [], {
            "zip_filename": name,
            "error": "ZIP_MISSING",
        }

    try:
        with ZipFile(path) as archive:
            for member in archive.namelist():
                if member.endswith("/"):
                    continue

                lower = member.lower()

                if not lower.endswith(
                    (".pdf", ".xls", ".xlsx")
                ):
                    continue

                try:
                    data = archive.read(member)
                    filename = Path(member).name
                    digest = hashlib.sha256(
                        data
                    ).hexdigest()

                    if (
                        digest[:16]
                        in SHARED_HASH_PREFIXES
                        or filename.lower()
                        in GENERIC_DOCUMENTS
                    ):
                        skipped_shared += 1
                        continue

                    if lower.endswith(".pdf"):
                        scanned_pdfs += 1

                        with pymupdf.open(
                            stream=data,
                            filetype="pdf",
                        ) as pdf:
                            for page_index in range(
                                min(
                                    len(pdf),
                                    MAX_PDF_PAGES,
                                )
                            ):
                                page_text = pdf[
                                    page_index
                                ].get_text()

                                candidates.extend(
                                    scan_text(
                                        name,
                                        base,
                                        filename,
                                        f"PDF page {page_index + 1}",
                                        page_text,
                                    )
                                )

                    else:
                        scanned_boqs += 1

                        for (
                            sheet_name,
                            row_number,
                            cells,
                        ) in spreadsheet_rows(
                            data, filename
                        ):
                            # Keep columns separate:
                            # label and amount must occur
                            # in the same or adjacent cell.
                            for col_index, cell in enumerate(
                                cells
                            ):
                                if not any(
                                    pattern.search(cell)
                                    for _, _, pattern
                                    in LABELS
                                ):
                                    continue

                                segment = cell

                                if col_index + 1 < len(
                                    cells
                                ):
                                    adjacent = cells[
                                        col_index + 1
                                    ]
                                    if not any(
                                        p.search(adjacent)
                                        for _, _, p in LABELS
                                    ):
                                        segment += (
                                            "\n" + adjacent
                                        )

                                candidates.extend(
                                    scan_text(
                                        name,
                                        base,
                                        filename,
                                        (
                                            f"{sheet_name} "
                                            f"row {row_number} "
                                            f"col {col_index + 1}"
                                        ),
                                        segment,
                                    )
                                )

                except Exception as exc:
                    errors.append(
                        f"{member}: {exc}"
                    )

    except Exception as exc:
        errors.append(f"ZIP: {exc}")

    unique = {}

    for candidate in candidates:
        key = (
            candidate["document"],
            candidate["location"],
            candidate["candidate_field"],
            candidate["candidate_value"],
        )
        unique.setdefault(key, candidate)

    result = list(unique.values())

    return result, {
        "zip_filename": name,
        "pdfs_scanned": scanned_pdfs,
        "boqs_scanned": scanned_boqs,
        "shared_documents_skipped": skipped_shared,
        "candidate_count": len(result),
        "errors": errors,
    }


def main():
    if not INPUT.exists():
        raise FileNotFoundError(
            f"Run V8.1 first: {INPUT.name}"
        )

    with INPUT.open(
        newline="",
        encoding="utf-8-sig",
    ) as handle:
        metadata = list(
            csv.DictReader(handle)
        )

    all_candidates = []
    records = []

    print("=" * 65)
    print("JAS V9 — FINANCIAL EVIDENCE AUDIT")
    print("=" * 65)

    for index, row in enumerate(metadata, 1):
        candidates, record = inspect_zip(row)

        all_candidates.extend(candidates)
        records.append(record)

        print(
            f"[{index:02d}/{len(metadata)}] "
            f"{row['zip_filename']} | "
            f"Candidates: {len(candidates)} | "
            f"PDFs: {record.get('pdfs_scanned', 0)} | "
            f"BOQs: {record.get('boqs_scanned', 0)}"
        )

    with OUTPUT.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=FIELDS,
        )
        writer.writeheader()
        writer.writerows(all_candidates)

    audit = {
        "version": "v9",
        "created_at": datetime.now(
            timezone.utc
        ).isoformat(),
        "tender_count": len(metadata),
        "candidate_count": len(all_candidates),
        "field_counts": dict(
            Counter(
                c["candidate_field"]
                for c in all_candidates
            )
        ),
        "review_status_counts": dict(
            Counter(
                c["review_status"]
                for c in all_candidates
            )
        ),
        "records": records,
    }

    AUDIT.write_text(
        json.dumps(
            audit,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    print("\nV9 SUMMARY")
    print("Tender ZIPs:", len(metadata))
    print("Candidates:", len(all_candidates))
    print("By field:", audit["field_counts"])
    print(
        "By review status:",
        audit["review_status_counts"],
    )
    print("\nSaved:", OUTPUT.name)
    print("Saved:", AUDIT.name)
    print("\nNo financial values were promoted.")


if __name__ == "__main__":
    main()
