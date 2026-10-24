
"""
JAS Tender Scraper — V8 BOQ Metadata Recovery

Reads .xls/.xlsx BOQs directly from tender ZIPs.
Uses V5 metadata as the baseline.

Does NOT automatically promote BOQ-derived titles,
financial values, or tender IDs.

Outputs:
  tender_boq_candidates_v8.csv
  tender_metadata_batch_v8.csv
  tender_boq_audit_v8.json
"""

import csv
import io
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

import xlrd
from openpyxl import load_workbook

import tender_pdf_extractor_v5 as v5


ROOT = Path(__file__).resolve().parent
ZIP_DIR = ROOT / "downloaded_tenders"

CANDIDATE_OUTPUT = ROOT / "tender_boq_candidates_v8.csv"
METADATA_OUTPUT = ROOT / "tender_metadata_batch_v8.csv"
AUDIT_OUTPUT = ROOT / "tender_boq_audit_v8.json"

MAX_ROWS_PER_SHEET = 150
MAX_COLUMNS = 25

CANDIDATE_FIELDS = [
    "zip_filename",
    "tender_id",
    "tender_id_status",
    "boq_filename",
    "sheet",
    "row_number",
    "candidate_field",
    "candidate_value",
    "candidate_type",
    "review_status",
    "evidence",
]

EXTRA_FIELDS = [
    "v8_boq_count",
    "v8_candidate_count",
    "v8_title_candidate_count",
    "v8_review_status",
    "v8_errors_json",
]

TITLE_LABEL = re.compile(
    r"\b(?:name\s+of\s+(?:the\s+)?work|"
    r"title\s+of\s+(?:the\s+)?work|"
    r"description\s+of\s+(?:the\s+)?work|"
    r"work\s+description|"
    r"name\s+of\s+package|"
    r"tender\s+description|"
    r"subject\s+of\s+tender)\b",
    re.I,
)

GENERIC_LABEL = re.compile(
    r"^(?:item\s*(?:no|number)?|"
    r"serial\s*(?:no|number)?|"
    r"description|quantity|unit|"
    r"rate|amount|total|"
    r"basic\s+rate|"
    r"bidder\s+name|"
    r"schedule\s+of\s+quantities|"
    r"bill\s+of\s+quantities|"
    r"boq|price\s+bid|"
    r"quoted\s+rate)$",
    re.I,
)

REJECT_TITLE = re.compile(
    r"contractual\s+cost|"
    r"funding\s+of\s+project|"
    r"name\s+of\s+contractor|"
    r"general\s+conditions|"
    r"terms\s+and\s+conditions|"
    r"do\s+not\s+modify|"
    r"protected\s+sheet|"
    r"quoted\s+rate|"
    r"percentage\s+of\s+amount",
    re.I,
)

ITEM_HEADER = re.compile(
    r"^(?:item\s*(?:no|number)?|"
    r"sl\.?\s*no\.?|"
    r"serial\s*(?:no|number)?|"
    r"item\s+description|"
    r"description\s+of\s+items?|"
    r"quantity|unit|rate|amount)$",
    re.I,
)


def norm(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def plausible_title(value):
    value = norm(value)

    if not 25 <= len(value) <= 350:
        return False

    if GENERIC_LABEL.fullmatch(value):
        return False

    if REJECT_TITLE.search(value):
        return False

    if re.fullmatch(r"[\d\s.,:/()%+-]+", value):
        return False

    return bool(re.search(r"[A-Za-z]{3,}", value))


def load_sheets(data, filename):
    """
    Return [(sheet_name, rows)] for XLS/XLSX.
    """
    if filename.lower().endswith(".xlsx"):
        workbook = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
        )
        try:
            result = []
            for sheet in workbook.worksheets:
                rows = [
                    list(row)
                    for row in sheet.iter_rows(
                        min_row=1,
                        max_row=min(
                            sheet.max_row or MAX_ROWS_PER_SHEET,
                            MAX_ROWS_PER_SHEET,
                        ),
                        max_col=min(
                            sheet.max_column or MAX_COLUMNS,
                            MAX_COLUMNS,
                        ),
                        values_only=True,
                    )
                ]
                result.append((sheet.title, rows))
            return result
        finally:
            workbook.close()

    workbook = xlrd.open_workbook(
        file_contents=data,
        on_demand=True,
    )

    try:
        result = []
        for sheet in workbook.sheets():
            rows = [
                [
                    sheet.cell_value(r, c)
                    for c in range(
                        min(sheet.ncols, MAX_COLUMNS)
                    )
                ]
                for r in range(
                    min(sheet.nrows, MAX_ROWS_PER_SHEET)
                )
            ]
            result.append((sheet.name, rows))
        return result
    finally:
        workbook.release_resources()


def make_candidate(
    zip_name,
    base_row,
    boq_filename,
    sheet_name,
    row_number,
    value,
    candidate_type,
    evidence,
):
    return {
        "zip_filename": zip_name,
        "tender_id": base_row.get("tender_id", ""),
        "tender_id_status": base_row.get(
            "tender_id_status", ""
        ),
        "boq_filename": boq_filename,
        "sheet": sheet_name,
        "row_number": row_number,
        "candidate_field": "tender_title",
        "candidate_value": norm(value),
        "candidate_type": candidate_type,
        "review_status": "NEEDS_REVIEW",
        "evidence": norm(evidence)[:500],
    }


def extract_sheet_candidates(
    zip_name,
    base_row,
    boq_filename,
    sheet_name,
    rows,
):
    """
    Recover explicitly labelled work titles.

    Item descriptions are not treated as tender titles.
    """
    candidates = []

    # Look for work-level titles in the header area.
    for row_index, row in enumerate(rows[:60]):
        cells = [norm(cell) for cell in row]

        for col_index, cell in enumerate(cells):
            match = TITLE_LABEL.search(cell)
            if not match:
                continue

            # Case 1: "Name of Work: Repair of roads"
            inline = norm(
                cell[match.end():].lstrip(" :-")
            )

            possible = []

            if inline:
                possible.append(
                    ("explicit_label_inline", inline)
                )

            # Case 2: label in one cell, value next to it.
            for next_cell in cells[col_index + 1:col_index + 4]:
                if next_cell:
                    possible.append(
                        ("explicit_label_adjacent", next_cell)
                    )
                    break

            # Case 3: value on the following row.
            if row_index + 1 < len(rows):
                following = [
                    norm(value)
                    for value in rows[row_index + 1]
                ]

                for next_cell in following[
                    col_index:col_index + 3
                ]:
                    if next_cell and not TITLE_LABEL.search(
                        next_cell
                    ):
                        possible.append(
                            ("explicit_label_next_row", next_cell)
                        )
                        break

            for candidate_type, value in possible:
                if not plausible_title(value):
                    continue

                candidates.append(
                    make_candidate(
                        zip_name,
                        base_row,
                        boq_filename,
                        sheet_name,
                        row_index + 1,
                        value,
                        candidate_type,
                        " | ".join(
                            cell for cell in cells if cell
                        ),
                    )
                )

                # Prefer the first plausible value
                # associated with each explicit label.
                break

    # Deduplicate within the sheet.
    unique = {}
    for candidate in candidates:
        key = candidate["candidate_value"].casefold()
        unique.setdefault(key, candidate)

    return list(unique.values())


def inspect_zip(zip_name, base_row):
    path = ZIP_DIR / zip_name
    candidates = []
    errors = []
    boq_count = 0

    if not path.exists():
        errors.append("ZIP not found")
        return candidates, boq_count, errors

    try:
        with ZipFile(path) as archive:
            members = [
                member
                for member in archive.namelist()
                if member.lower().endswith(
                    (".xls", ".xlsx")
                )
                and not member.endswith("/")
            ]

            for member in sorted(members):
                boq_count += 1
                filename = Path(member).name

                try:
                    sheets = load_sheets(
                        archive.read(member),
                        filename,
                    )

                    for sheet_name, rows in sheets:
                        candidates.extend(
                            extract_sheet_candidates(
                                zip_name,
                                base_row,
                                filename,
                                sheet_name,
                                rows,
                            )
                        )

                except Exception as exc:
                    errors.append(
                        f"{filename}: {exc}"
                    )

    except Exception as exc:
        errors.append(f"ZIP error: {exc}")

    unique = {}

    for candidate in candidates:
        key = (
            candidate["candidate_field"],
            candidate["candidate_value"].casefold(),
        )
        unique.setdefault(key, candidate)

    return list(unique.values()), boq_count, errors


def main():
    mappings = v5.load_mappings()

    metadata_rows = []
    review_rows = []
    audit_records = []

    print("=" * 68)
    print("JAS TENDER METADATA — V8 BOQ RECOVERY")
    print("=" * 68)

    for index, zip_name in enumerate(v5.BATCH, 1):
        base_row = v5.process_zip(
            zip_name,
            mappings,
        )

        candidates, boq_count, errors = inspect_zip(
            zip_name,
            base_row,
        )

        # Keep existing V5 values unchanged.
        # Only missing titles are actionable.
        for candidate in candidates:
            if base_row.get("tender_title"):
                candidate["review_status"] = (
                    "ALREADY_HAS_V5_TITLE"
                )

        actionable = [
            candidate
            for candidate in candidates
            if candidate["review_status"] == "NEEDS_REVIEW"
        ]

        row = dict(base_row)
        row.update({
            "v8_boq_count": boq_count,
            "v8_candidate_count": len(candidates),
            "v8_title_candidate_count": len(actionable),
            "v8_review_status": (
                "NEEDS_REVIEW" if actionable
                else "EXTRACTION_ERROR" if errors
                else "NO_NEW_TITLE_CANDIDATES"
            ),
            "v8_errors_json": json.dumps(
                errors,
                ensure_ascii=False,
            ),
        })

        metadata_rows.append(row)
        review_rows.extend(candidates)

        audit_records.append({
            "zip_filename": zip_name,
            "tender_id": base_row["tender_id"],
            "tender_id_status": base_row[
                "tender_id_status"
            ],
            "boq_count": boq_count,
            "candidates": candidates,
            "errors": errors,
        })

        print(
            f"[{index:02d}/{len(v5.BATCH)}] "
            f"{zip_name} | "
            f"BOQs: {boq_count} | "
            f"Candidates: {len(candidates)} | "
            f"New titles: {len(actionable)}"
        )

    with CANDIDATE_OUTPUT.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=CANDIDATE_FIELDS,
        )
        writer.writeheader()
        writer.writerows(review_rows)

    with METADATA_OUTPUT.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=v5.FIELDS + EXTRA_FIELDS,
        )
        writer.writeheader()
        writer.writerows(metadata_rows)

    audit = {
        "version": "v8",
        "created_at": datetime.now(
            timezone.utc
        ).isoformat(),
        "tenders": len(metadata_rows),
        "boq_files": sum(
            row["v8_boq_count"]
            for row in metadata_rows
        ),
        "total_candidates": len(review_rows),
        "actionable_title_candidates": sum(
            row["v8_title_candidate_count"]
            for row in metadata_rows
        ),
        "status_counts": dict(
            Counter(
                row["v8_review_status"]
                for row in metadata_rows
            )
        ),
        "records": audit_records,
    }

    AUDIT_OUTPUT.write_text(
        json.dumps(
            audit,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    print("\nV8 SUMMARY")
    print("Tender ZIPs:", len(metadata_rows))
    print("BOQ files:", audit["boq_files"])
    print("Candidates:", audit["total_candidates"])
    print(
        "Actionable title candidates:",
        audit["actionable_title_candidates"],
    )

    print("\nSaved:", CANDIDATE_OUTPUT.name)
    print("Saved:", METADATA_OUTPUT.name)
    print("Saved:", AUDIT_OUTPUT.name)

    print("\nV5 and V7 outputs remain unchanged.")


if __name__ == "__main__":
    main()
